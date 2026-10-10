"""Gera o manual em DOCX (para editar no Google Docs) a partir do manual.html, com as mesmas telas.
Uso: python3 manual/capturas/gerar_docx.py SAIDA.docx
Cada seção começa depois de uma quebra de página explícita — no Google Docs ela aparece como uma linha
"quebra de página" e dá para apagar ou mover. O DOCX não vai para o repositório: é entregue ao dono."""
import os, re, sys, tempfile
import lxml.html
import segno
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

RAIZ = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
SAIDA = sys.argv[1] if len(sys.argv) > 1 else os.path.join(RAIZ, "manual.docx")

INK = RGBColor(0x0F, 0x17, 0x2A); INK2 = RGBColor(0x33, 0x41, 0x55); INK3 = RGBColor(0x64, 0x74, 0x8B)
AZUL = RGBColor(0x25, 0x63, 0xA8)
CAIXA = {  # classe: (fundo, cor da borda esquerda)
    "dica": ("EEF4FB", "2563A8"), "nota": ("F3F4F6", "94A3B8"),
    "aviso": ("FDF3E1", "D97706"), "fala": ("F1F5F9", "2563A8")}
LARGURA = {None: 16.5, "ficha": 11.0, "ficha-t": 9.0, "cel": 6.2, "telefone": 8.0}

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)

st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10.5); st.font.color.rgb = INK
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.15
for nome, tam, cor in (("Heading 1", 18, INK), ("Heading 2", 13, INK), ("Title", 34, INK)):
    h = doc.styles[nome]; h.font.name = "Arial"; h.font.size = Pt(tam); h.font.bold = True; h.font.color.rgb = cor
    rpr = h.element.get_or_add_rPr(); f = rpr.find(qn("w:rFonts"))
    if f is None: f = OxmlElement("w:rFonts"); rpr.append(f)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"): f.set(qn(a), "Arial")
doc.styles["Heading 1"].paragraph_format.space_before = Pt(0); doc.styles["Heading 1"].paragraph_format.space_after = Pt(10)
doc.styles["Heading 2"].paragraph_format.space_before = Pt(14); doc.styles["Heading 2"].paragraph_format.space_after = Pt(4)

# ---------- utilidades ----------
def limpo(t):
    return re.sub(r"\s+", " ", t or "")

def hyperlink(par, url, texto, negrito=False):
    rid = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "2563A8"); rpr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    if negrito: rpr.append(OxmlElement("w:b"))
    r.append(rpr); t = OxmlElement("w:t"); t.text = texto; t.set(qn("xml:space"), "preserve"); r.append(t)
    h.append(r); par._p.append(h)

def inline(par, el, b=False, i=False, mono=False, tam=None, cor=None, pular_primeiro_b=False):
    """Escreve o texto de el (com filhos) no parágrafo, respeitando negrito/itálico/link/quebra."""
    def run(txt, b, i, mono):
        txt = limpo(txt)
        if not txt: return
        r = par.add_run(txt); r.bold = b or None; r.italic = i or None
        if mono: r.font.name = "Courier New"
        if tam: r.font.size = Pt(tam)
        if cor is not None: r.font.color.rgb = cor
    run(el.text, b, i, mono)
    pulou = False
    for c in el:
        if not isinstance(c.tag, str):
            run(c.tail, b, i, mono); continue
        tag = c.tag
        if pular_primeiro_b and not pulou and tag in ("b", "strong"):
            pulou = True; run(c.tail, b, i, mono); continue
        if tag == "br":
            par.add_run().add_break()
        elif tag in ("b", "strong"):
            inline(par, c, True, i, mono, tam, cor)
        elif tag in ("i", "em"):
            inline(par, c, b, True, mono, tam, cor)
        elif tag in ("code", "kbd"):
            inline(par, c, b, i, True, tam, cor)
        elif tag == "a":
            href = c.get("href") or ""
            if href.startswith("http"): hyperlink(par, href, limpo(c.text_content()), b)
            else: inline(par, c, b, i, mono, tam, cor)
        elif tag in ("ul", "ol", "figure", "table", "div", "svg", "img"):
            pass  # blocos: tratados fora
        else:
            inline(par, c, b, i, mono, tam, cor)
        run(c.tail, b, i, mono)

def sombrear(par, fundo, borda):
    ppr = par._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    e = OxmlElement("w:left")   # só a barra colorida à esquerda e o fundo: o Google Docs importa sem linhas entre os parágrafos
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "18"); e.set(qn("w:space"), "8"); e.set(qn("w:color"), borda)
    bd.append(e)
    ppr.append(bd)
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fundo)
    ppr.append(shd)
    par.paragraph_format.left_indent = Cm(0.35); par.paragraph_format.right_indent = Cm(0.2)

def quebra():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def lista(el, numerada=False, caixa=None, nivel=0):
    n = 0
    for li in el:
        if li.tag != "li": continue
        n += 1
        p = doc.add_paragraph()
        pf = p.paragraph_format; pf.left_indent = Cm(0.9 + 0.6 * nivel); pf.first_line_indent = Cm(-0.5); pf.space_after = Pt(3)
        p.add_run(("%d.  " % n) if numerada else "•  ").bold = numerada or None
        inline(p, li)
        if caixa: sombrear(p, *caixa); pf.left_indent = Cm(1.0 + 0.6 * nivel)
        for sub in li:
            if sub.tag in ("ul", "ol"): lista(sub, sub.tag == "ol", caixa, nivel + 1)

def tabela(el):
    linhas = [tr for tr in el.iter("tr")]
    ncol = max(len([c for c in tr if c.tag in ("td", "th")]) for tr in linhas)
    t = doc.add_table(rows=len(linhas), cols=ncol); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    total = 17.0; larg = [4.2] + [(total - 4.2) / (ncol - 1)] * (ncol - 1) if ncol > 1 else [total]
    for r, tr in enumerate(linhas):
        cels = [c for c in tr if c.tag in ("td", "th")]
        for k, c in enumerate(cels):
            cell = t.cell(r, k); cell.width = Cm(larg[k])
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
            inline(p, c, b=(c.tag == "th"), tam=9.5)
            if c.tag == "th":
                tcpr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "EEF2F7"); tcpr.append(shd)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def figura(el):
    img = el.find(".//img")
    if img is None: return
    caminho = os.path.join(RAIZ, img.get("src").split("?")[0])
    if not os.path.exists(caminho): return
    w = LARGURA.get(el.get("class"), LARGURA[None])
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(caminho, width=Cm(w))
    cap = el.find("figcaption")
    if cap is not None:
        c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(10)
        inline(c, cap, tam=9, cor=INK3, i=False)

def caixa(el):
    """Caixa colorida (dica, nota, aviso, fala): título, texto e listas na ordem em que aparecem."""
    fundo, borda = CAIXA.get(el.get("class"), CAIXA["nota"])
    cls = el.get("class"); filhos = [c for c in el if isinstance(c.tag, str)]
    tem_titulo = bool(filhos) and filhos[0].tag in ("b", "strong") and not (el.text or "").strip() and cls != "fala"
    def novo():
        q = doc.add_paragraph(); sombrear(q, fundo, borda); q.paragraph_format.space_after = Pt(0); return q
    def fecha(q):
        if not q.text.strip() and not q._p.findall(qn("w:r") + "/" + qn("w:drawing")): q._p.getparent().remove(q._p)
    if tem_titulo:
        q = novo(); r = q.add_run(limpo(filhos[0].text_content())); r.bold = True
    q = novo(); italico = cls == "fala"
    fake = lxml.html.fragment_fromstring("<span></span>")
    def escreve(txt):
        if txt and limpo(txt).strip():
            fake.text = txt; inline(q, fake, i=italico)
    escreve(el.text)
    for k, c in enumerate(el):
        if not isinstance(c.tag, str): escreve(c.tail); continue
        if k == 0 and tem_titulo: escreve(c.tail); continue
        if c.tag in ("ul", "ol", "table"):
            fecha(q)
            if c.tag == "table": tabela(c)
            else: lista(c, c.tag == "ol", (fundo, borda))
            q = novo(); escreve(c.tail)
        else:
            casca = lxml.html.fragment_fromstring("<span></span>"); casca.append(lxml.html.fragment_fromstring(lxml.html.tostring(c, encoding="unicode", with_tail=False)))
            inline(q, casca, i=italico); escreve(c.tail)
    fecha(q)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

# ---------- capa ----------
html = lxml.html.parse(os.path.join(RAIZ, "manual.html")).getroot()
pag = html.find_class("pagina")[0]
capa = pag.find_class("capa")[0]
tmp = tempfile.mkdtemp()

p = doc.add_paragraph(); p.add_run().add_picture(os.path.join(RAIZ, "icons", "icon-192.png"), width=Cm(2.4))
p = doc.add_paragraph(); r = p.add_run("Caderneta"); r.bold = True; r.font.size = Pt(34); r.font.name = "Georgia"
p.paragraph_format.space_after = Pt(0)
p = doc.add_paragraph(); r = p.add_run("SUAS CONTAS, DO SEU JEITO"); r.font.size = Pt(9); r.font.color.rgb = INK3
p.paragraph_format.space_after = Pt(70)
p = doc.add_paragraph(); r = p.add_run(limpo(capa.find_class("titulo")[0].text_content())); r.bold = True; r.font.size = Pt(26)
p = doc.add_paragraph(); r = p.add_run(limpo(capa.find_class("sub")[0].text_content())); r.font.size = Pt(12); r.font.color.rgb = INK2
p = doc.add_paragraph(); r = p.add_run(limpo(capa.find_class("edicao")[0].text_content())); r.font.size = Pt(9.5); r.font.name = "Courier New"; r.font.color.rgb = INK2
p.paragraph_format.space_after = Pt(90)
quem = capa.find_class("quem")[0]
p = doc.add_paragraph(); r = p.add_run("DESENVOLVIMENTO"); r.font.size = Pt(8); r.font.color.rgb = INK3; p.paragraph_format.space_after = Pt(0)
p = doc.add_paragraph(); r = p.add_run(limpo(quem.find_class("n")[0].text_content())); r.bold = True; r.font.size = Pt(13); p.paragraph_format.space_after = Pt(0)
p = doc.add_paragraph(); r = p.add_run(limpo(quem.find_class("c")[0].text_content())); r.font.size = Pt(9.5); r.font.color.rgb = INK2
for a in capa.find_class("links")[0].iter("a"):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1)
    rotulo = limpo(a.getparent().text or "")
    p.add_run(rotulo).font.size = Pt(9.5); hyperlink(p, a.get("href"), limpo(a.text_content()))
# os dois QR, lado a lado
t = doc.add_table(rows=2, cols=2); t.alignment = WD_TABLE_ALIGNMENT.LEFT
for k, (url, rot) in enumerate((("https://rgbittencourt.github.io/caderneta/", "app"), ("https://github.com/rgbittencourt/caderneta", "código"))):
    arq = os.path.join(tmp, "qr%d.png" % k); segno.make(url, error="m").save(arq, scale=8, border=2, dark="#0F172A")
    c = t.cell(0, k); c.width = Cm(3.2); c.paragraphs[0].add_run().add_picture(arq, width=Cm(2.6))
    c2 = t.cell(1, k); c2.width = Cm(3.2); rr = c2.paragraphs[0].add_run(rot); rr.font.size = Pt(8); rr.font.color.rgb = INK3
quebra()

# ---------- sumário e o resto, na ordem do manual ----------
primeiro_h2 = True
for el in pag:
    if not isinstance(el.tag, str): continue
    tag, cls = el.tag, el.get("class")
    if tag == "header": continue
    if tag == "nav":
        doc.add_paragraph("Sumário", style="Heading 1")
        for k, li in enumerate(el.iter("li"), 1):
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1)
            p.add_run("%d.  " % k).bold = True; p.add_run(limpo(li.text_content()))
        doc.add_paragraph()
        continue
    if tag == "h2":
        quebra()
        doc.add_paragraph(limpo(el.text_content()), style="Heading 1")
    elif tag == "h3":
        doc.add_paragraph(limpo(el.text_content()), style="Heading 2")
    elif tag == "p":
        p = doc.add_paragraph(); inline(p, el)
    elif tag in ("ul", "ol"):
        lista(el, tag == "ol"); doc.add_paragraph().paragraph_format.space_after = Pt(0)
    elif tag == "table":
        tabela(el)
    elif tag == "figure":
        figura(el)
    elif tag == "div" and cls in CAIXA:
        caixa(el)
    elif tag == "div" and cls == "colofao":
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(24)
        p.add_run().add_picture(os.path.join(RAIZ, "icons", "icon-192.png"), width=Cm(1.4))
        b = el.find(".//b"); spans = el.findall(".//span")
        p = doc.add_paragraph(); r = p.add_run(limpo(b.text_content())); r.bold = True; r.font.size = Pt(13); p.paragraph_format.space_after = Pt(0)
        for s in spans:
            p = doc.add_paragraph(); r = p.add_run(limpo(s.text_content())); r.font.size = Pt(9); r.font.color.rgb = INK2; p.paragraph_format.space_after = Pt(0)
    elif tag == "div" and cls == "rodape":
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12)
        r = p.add_run(limpo(el.text_content())); r.font.size = Pt(9); r.font.color.rgb = INK3

# rodapé com o número da página
rod = sec.footer.paragraphs[0]; rod.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = rod.add_run("Caderneta · Manual de uso · página "); r.font.size = Pt(8); r.font.color.rgb = INK3
for tipo, txt in (("begin", None), ("instr", " PAGE "), ("end", None)):
    r = rod.add_run(); r.font.size = Pt(8); r.font.color.rgb = INK3
    if tipo == "instr":
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._r.append(it)
    else:
        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tipo); r._r.append(fc)

doc.core_properties.title = "Caderneta — Manual de uso"
doc.core_properties.author = "Caderneta"
doc.save(SAIDA)
print("salvo:", SAIDA)
